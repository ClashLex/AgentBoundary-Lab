"""Build the editable Unreliable Orchestra how-to-play presentation."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from random import Random

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "presentations" / "Unreliable_Orchestra_How_to_Play.pptx"
W, H = 13.333, 7.5

BG = "F8F7F3"
INK = "191919"
MUTED = "575757"
RULE = "CFD0C8"
WHITE = "FFFFFF"
LIME = "D5FA58"
LIME_SOFT = "E8F7B5"
LAVENDER = "DDD0F8"
LAVENDER_SOFT = "F3EDFC"
PEACH = "FFD9C7"
PEACH_SOFT = "FFF1E9"
REST = "EDEDEB"
MISS = "FFE4DC"
PURPLE = "9274CC"

HEAD = "Bahnschrift"
BODY = "Arial"


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def box(slide, x, y, w, h, fill=WHITE, stroke=INK, stroke_pt=1.2, shape=MSO_SHAPE.RECTANGLE):
    item = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = color(fill)
    if stroke is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = color(stroke)
        item.line.width = Pt(stroke_pt)
    return item


def text(slide, value, x, y, w, h, size=17, fill=INK, bold=False, font=BODY,
         align=PP_ALIGN.LEFT, middle=False, margin=0, italic=False):
    item = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = item.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(0.01)
    frame.margin_bottom = 0
    if middle:
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(value.split("\n")):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.alignment = align
        para.space_before = Pt(0)
        para.space_after = Pt(0)
        para.line_spacing = 1.04
        run = para.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color(fill)
    return item


def label(slide, value, x, y, w, h=0.24, fill=MUTED, size=10):
    return text(slide, value.upper(), x, y, w, h, size, fill, True, BODY)


def rule(slide, x, y, w, fill=RULE, h=0.012):
    return box(slide, x, y, w, h, fill, None)


def pill(slide, value, x, y, w, fill=WHITE, h=0.38, size=11):
    box(slide, x, y, w, h, fill, INK, 1.0)
    text(slide, value, x + 0.08, y, w - 0.16, h, size, INK, True, BODY,
         PP_ALIGN.CENTER, True)


def shadow_card(slide, x, y, w, h, fill=WHITE, shadow=LIME):
    box(slide, x + 0.10, y + 0.10, w, h, shadow, None)
    return box(slide, x, y, w, h, fill, INK, 1.3)


def logo(slide):
    box(slide, 0.63, 0.27, 0.42, 0.42, BG, INK, 1.5, MSO_SHAPE.OVAL)
    for i, ht in enumerate((0.14, 0.26, 0.18)):
        box(slide, 0.755 + i * 0.075, 0.38 + (0.26 - ht) / 2,
            0.038, ht, INK, None)
    text(slide, "UNRELIABLE\nORCHESTRA", 1.17, 0.24, 2.1, 0.46,
         10.5, INK, True, HEAD)


def background_image() -> BytesIO:
    # Small, pale ordered dots echo the website's restrained WebGL dither.
    image = Image.new("RGB", (1600, 900), color(BG))
    draw = ImageDraw.Draw(image)
    bayer = ((0, 8, 2, 10), (12, 4, 14, 6), (3, 11, 1, 9), (15, 7, 13, 5))
    rng = Random(401)
    for gy, py in enumerate(range(8, 900, 14)):
        for gx, px in enumerate(range(8, 1600, 14)):
            lavender_weight = max(0, 1 - ((px - 1450) / 830) ** 2 - ((py - 110) / 600) ** 2)
            lime_weight = max(0, 1 - ((px - 50) / 750) ** 2 - ((py - 890) / 620) ** 2)
            strength = max(lavender_weight, lime_weight)
            if strength < 0.17 or bayer[gy % 4][gx % 4] > strength * 12 + rng.random():
                continue
            ink = (232, 224, 245) if lavender_weight >= lime_weight else (225, 239, 197)
            draw.rectangle((px, py, px + 2, py + 2), fill=ink)
    output = BytesIO()
    image.save(output, format="PNG")
    output.seek(0)
    return output


PATTERN = background_image().getvalue()
prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
prs.core_properties.title = "Unreliable Orchestra — How to Play"
prs.core_properties.subject = "Beginner guide to the three-player cooperative music puzzle"
prs.core_properties.author = "Unreliable Orchestra"


def slide(number, title=None, subtitle=None, kicker="HOW TO PLAY"):
    page = prs.slides.add_slide(prs.slide_layouts[6])
    page.shapes.add_picture(BytesIO(PATTERN), 0, 0, Inches(W), Inches(H))
    logo(page)
    label(page, kicker, 10.66, 0.37, 1.72, fill=INK, size=10)
    text(page, f"{number:02d} / 12", 12.0, 0.33, 0.7, 0.28,
         10, MUTED, True, BODY, PP_ALIGN.RIGHT)
    rule(page, 0.63, 0.83, 12.07)
    rule(page, 0.63, 7.12, 12.07)
    label(page, "THREE SIGNALS. ONE SONG.", 0.67, 7.20, 3.1, 0.16, MUTED, 8)
    label(page, "MADE FOR PLAY", 11.15, 7.20, 1.52, 0.16, MUTED, 8)
    if title:
        text(page, title, 0.72, 1.06, 11.95, 0.58, 30, INK, True, HEAD)
    if subtitle:
        text(page, subtitle, 0.75, 1.68, 11.7, 0.54, 15, MUTED)
    return page


def slot_board(page, values, x, y, w, h, captions=False, numbers=True,
               font_size=24, miss_indices=()):
    gap = 0
    sw = (w - gap * 7) / 8
    for i, value in enumerate(values):
        sx = x + i * (sw + gap)
        is_rest = value in ("REST", "—", "R")
        is_unknown = value in ("?", "")
        bg = MISS if i in miss_indices else REST if is_rest else WHITE if is_unknown else LIME_SOFT
        box(page, sx, y, sw, h, bg, INK, 1.2)
        if numbers:
            label(page, f"{i + 1:02d}", sx + 0.09, y + 0.12, sw - 0.18, 0.20,
                  MUTED, 9)
        display = "—" if is_rest else value
        text(page, display, sx + 0.05, y + h * 0.26, sw - 0.10, h * 0.46,
             font_size, INK, True, HEAD, PP_ALIGN.CENTER, True)
        if captions:
            cap = "REST" if is_rest else "UNANSWERED" if is_unknown else "NOTE"
            label(page, cap, sx + 0.04, y + h - 0.28, sw - 0.08, 0.18,
                  MUTED, 7.3)


def token_row(page, values, x, y, width, fill=LAVENDER, value_size=22):
    gap = 0.14
    tw = (width - gap * (len(values) - 1)) / len(values)
    for i, value in enumerate(values):
        bx = x + i * (tw + gap)
        box(page, bx, y, tw, 0.72, fill, INK, 1.1)
        text(page, value, bx, y + 0.04, tw, 0.60, value_size,
             INK, True, HEAD, PP_ALIGN.CENTER, True)


# 01 — Invitation and visual identity
p = slide(1, kicker="A BEGINNER'S GUIDE")
label(p, "A COOPERATIVE MUSIC PUZZLE", 0.79, 1.35, 4.9, fill=MUTED, size=11)
text(p, "UNRELIABLE", 0.77, 1.80, 6.5, 0.79, 45, INK, True, HEAD)
text(p, "ORCHESTRA", 0.77, 2.60, 6.5, 0.80, 45, PURPLE, True, HEAD)
text(p, "Three people. Three clues.\nOne missing song.", 0.81, 3.75, 5.6, 1.20,
     23, INK, True, HEAD)
text(p, "Learn the rules in five minutes. No musical training needed.",
     0.81, 5.32, 5.55, 0.60, 16, MUTED)
pill(p, "3 PLAYERS", 0.81, 6.16, 1.28, LIME)
pill(p, "8 SLOTS", 2.22, 6.16, 1.12, WHITE)
pill(p, "1 SHARED ANSWER", 3.47, 6.16, 2.06, LAVENDER)
shadow_card(p, 7.10, 1.44, 5.19, 4.91, LAVENDER, LIME)
label(p, "THE MISSING COMPOSITION", 7.42, 1.76, 3.7, fill=INK, size=10)
for j in range(4):
    rule(p, 7.47, 2.39 + j * 0.43, 4.41, "9B8BB0", 0.018)
for cx, cy in ((8.09, 3.02), (9.41, 2.57), (10.85, 3.50)):
    box(p, cx, cy, 0.29, 0.21, INK, None, shape=MSO_SHAPE.OVAL)
    rule(p, cx + 0.26, cy - 0.56, 0.025, INK, 0.69)
box(p, 10.13, 2.49, 0.77, 0.77, LIME, INK, 1.5, MSO_SHAPE.OVAL)
text(p, "?", 10.16, 2.51, 0.70, 0.66, 36, INK, True, HEAD, PP_ALIGN.CENTER, True)
slot_board(p, ["?", "?", "?", "?", "?", "?", "?", "?"], 7.45, 5.09, 4.40, 0.80,
           numbers=False, font_size=19)

# 02 — Goal and core terms
p = slide(2, "The goal: find the exact eight-slot song.",
          "A hidden song is chosen for each round. Your team reconstructs it together.")
label(p, "THE HIDDEN TARGET", 0.80, 2.31, 3.1, fill=INK, size=11)
slot_board(p, ["?", "?", "?", "?", "?", "?", "?", "?"], 0.80, 2.64, 11.73, 1.58,
           captions=True, font_size=37)
for x, bg, term, desc in (
    (0.80, LAVENDER_SOFT, "TARGET", "The original eight-slot song. It stays hidden until the reveal."),
    (4.75, WHITE, "BOARD", "The answer everyone can see. The Conductor fills its eight slots."),
    (8.70, LIME_SOFT, "ROUND", "One puzzle, from the ready check through scoring and reveal."),
):
    box(p, x, 4.64, 3.83, 1.43, bg, INK, 1.1)
    label(p, term, x + 0.18, 4.87, 3.41, fill=INK, size=12)
    text(p, desc, x + 0.18, 5.21, 3.45, 0.69, 15, MUTED)
text(p, "A perfect answer matches every note AND every rest in its exact position.",
     0.82, 6.37, 11.7, 0.46, 18, INK, True, HEAD)

# 03 — Music terms
p = slide(3, "The music words are simple.",
          "You only need five pitches, silence, and eight numbered positions.")
for x, bg, term, mark, desc in (
    (0.80, LIME_SOFT, "NOTE / PITCH", "C", "A note makes sound. Its pitch is how high or low it sounds. This game uses C, D, E, G and A."),
    (4.75, REST, "REST", "—", "A rest is an intentional silent beat. It still counts as a filled slot and can earn a point."),
    (8.70, WHITE, "UNANSWERED", "?", "A question mark means the team has not chosen a note or rest yet. You cannot submit it."),
):
    box(p, x, 2.35, 3.83, 2.80, bg, INK, 1.2)
    label(p, term, x + 0.22, 2.60, 3.30, fill=INK, size=11)
    text(p, mark, x + 0.22, 3.04, 0.85, 0.88, 44, INK, True, HEAD)
    text(p, desc, x + 0.22, 4.05, 3.30, 0.91, 15, MUTED)
box(p, 0.80, 5.52, 11.73, 0.98, LAVENDER, INK, 1.2)
text(p, "SLOT", 1.03, 5.78, 1.03, 0.26, 13, INK, True, HEAD)
text(p, "One numbered beat in the song. At 100 BPM (beats per minute), a beat takes 0.6 seconds.",
     2.00, 5.73, 10.20, 0.52, 17, INK)

# 04 — Joining
p = slide(4, "Get three people into one room.",
          "Everyone opens the same game link. A room code brings the team together.")
steps = (
    ("01", "CREATE", "One player creates a room and becomes the host."),
    ("02", "SHARE", "The host sends the six-character room code to two friends."),
    ("03", "JOIN", "Each friend enters that code to join the same room."),
    ("04", "READY", "All three mark ready. The host starts the round."),
)
for i, (number, name, desc) in enumerate(steps):
    x = 0.80 + i * 3.00
    bg = (LAVENDER, WHITE, PEACH, LIME)[i]
    box(p, x, 2.55, 2.73, 2.55, bg, INK, 1.15)
    text(p, number, x + 0.17, 2.74, 0.67, 0.54, 27, INK, True, HEAD)
    label(p, name, x + 0.17, 3.43, 2.27, fill=INK, size=12)
    text(p, desc, x + 0.17, 3.79, 2.31, 1.08, 16, MUTED)
    if i < 3:
        text(p, "→", x + 2.78, 3.50, 0.24, 0.38, 20, PURPLE, True, HEAD)
box(p, 0.80, 5.54, 11.73, 0.89, WHITE, INK, 1.1)
text(p, "ROOM = shared game space     HOST = person who created it     READY = safe to begin",
     1.08, 5.84, 11.08, 0.30, 15, INK, True, HEAD)
text(p, "Three devices work remotely; three separate browser tabs work for a local demo.",
     0.82, 6.68, 11.69, 0.25, 13, MUTED)

# 05 — Roles
p = slide(5, "Each role holds a different piece.",
          "No one can solve the complete song alone.")
roles = (
    (0.80, LAVENDER, "CONDUCTOR", "Knows how many steps to move pitches backward.",
     "Places notes and rests on the shared board, then submits."),
    (4.75, LIME, "RHYTHM KEEPER", "Sees which slots contain a note and which are rests.",
     "Tells the team where the sounding notes belong."),
    (8.70, PEACH, "MELODIST", "Sees the shifted pitches in playing order.",
     "Shares the note sequence; rest positions are hidden."),
)
for x, bg, name, knows, job in roles:
    box(p, x, 2.38, 3.83, 3.88, bg, INK, 1.2)
    label(p, name, x + 0.22, 2.65, 3.35, fill=INK, size=13)
    rule(p, x + 0.22, 3.10, 3.35, INK)
    label(p, "PRIVATE CLUE", x + 0.22, 3.38, 3.35, fill=MUTED, size=9)
    text(p, knows, x + 0.22, 3.68, 3.28, 0.98, 17, INK, True, HEAD)
    label(p, "TEAM JOB", x + 0.22, 4.94, 3.35, fill=MUTED, size=9)
    text(p, job, x + 0.22, 5.23, 3.25, 0.82, 15, INK)
text(p, "The host is Conductor. The first joiner is Rhythm Keeper; the second is Melodist.",
     0.82, 6.59, 11.80, 0.30, 14, MUTED)

# 06 — Shift
p = slide(6, "What does a pitch shift mean?",
          "Imagine the five allowed notes arranged in a loop.")
for i, note in enumerate(("C", "D", "E", "G", "A")):
    x = 1.15 + i * 2.37
    box(p, x, 2.43, 0.99, 0.99, LAVENDER if i == 1 else WHITE,
        INK, 1.35, MSO_SHAPE.OVAL)
    text(p, note, x, 2.58, 0.99, 0.66, 31, INK, True, HEAD, PP_ALIGN.CENTER, True)
    if i < 4:
        text(p, "→", x + 1.36, 2.71, 0.42, 0.34, 27, PURPLE, True, HEAD)
text(p, "then back to C ↺", 9.41, 3.67, 2.75, 0.38, 16, PURPLE, True, HEAD,
     PP_ALIGN.CENTER)
box(p, 0.81, 4.25, 5.55, 1.58, LIME_SOFT, INK, 1.2)
label(p, "SHIFTED NOTE", 1.04, 4.48, 2.0, fill=INK)
text(p, "D", 1.04, 4.83, 0.85, 0.63, 31, INK, True, HEAD)
text(p, "1 step backward", 2.12, 4.97, 2.47, 0.32, 17, INK)
text(p, "→  C", 4.66, 4.88, 1.36, 0.47, 24, INK, True, HEAD)
box(p, 6.65, 4.25, 5.86, 1.58, LAVENDER_SOFT, INK, 1.2)
label(p, "HOW TO USE THE CLUE", 6.90, 4.49, 4.75, fill=INK)
text(p, "The Melodist sees the shifted note. The Conductor knows the number of steps to undo it.",
     6.90, 4.86, 5.32, 0.69, 16, INK)
text(p, "A shift changes pitch, not the notes' playing order. The Conductor's number is 1–4.",
     0.82, 6.33, 11.65, 0.38, 15, MUTED)

# 07 — Three clues from one worked example
p = slide(7, "See the three clues for one song.",
          "This teaching example uses a one-step shift. The hidden answer is still covered.")
box(p, 0.80, 2.38, 3.83, 3.78, LIME_SOFT, INK, 1.2)
label(p, "RHYTHM KEEPER", 1.03, 2.63, 3.18, fill=INK, size=11)
text(p, "Note / rest map", 1.03, 2.99, 3.27, 0.43, 20, INK, True, HEAD)
slot_board(p, ["N", "R", "N", "N", "R", "N", "N", "R"],
           1.03, 3.74, 3.38, 1.17, numbers=True, font_size=16)
text(p, "N = note; — = rest.\nPitches remain unknown.",
     1.03, 5.20, 3.28, 0.67, 15, MUTED)
box(p, 4.75, 2.38, 3.83, 3.78, PEACH_SOFT, INK, 1.2)
label(p, "MELODIST", 4.98, 2.63, 3.17, fill=INK, size=11)
text(p, "Shifted notes", 4.98, 2.99, 3.26, 0.43, 20, INK, True, HEAD)
token_row(p, ["D", "G", "A", "C", "E"], 4.98, 3.77, 3.37, PEACH, 18)
text(p, "1st to 5th sounding note.\nNo rest positions shown.",
     4.98, 5.20, 3.28, 0.67, 15, MUTED)
box(p, 8.70, 2.38, 3.83, 3.78, LAVENDER_SOFT, INK, 1.2)
label(p, "CONDUCTOR", 8.93, 2.63, 3.17, fill=INK, size=11)
text(p, "Undo the shift", 8.93, 2.99, 3.27, 0.43, 20, INK, True, HEAD)
box(p, 8.95, 3.78, 1.20, 1.17, LAVENDER, INK, 1.2)
text(p, "1", 8.95, 3.84, 1.20, 1.00, 50, INK, True, HEAD,
     PP_ALIGN.CENTER, True)
text(p, "step backward\nfor each note", 10.35, 3.99, 1.83, 0.83, 17,
     INK, True, HEAD)
text(p, "Cannot see the original\npitches or rests.", 8.93, 5.20, 3.27, 0.67,
     15, MUTED)
text(p, "They must tell one another what they know.", 0.83, 6.50, 11.59, 0.35,
     18, INK, True, HEAD)

# 08 — Solve the sample
p = slide(8, "Combine the clues to solve it.",
          "First decode the notes. Then place them in the note slots.")
label(p, "MELODIST'S SHIFTED PITCHES", 0.82, 2.30, 4.4, fill=INK)
token_row(p, ["D", "G", "A", "C", "E"], 0.82, 2.69, 5.65, PEACH, 23)
text(p, "→", 6.50, 2.88, 0.30, 0.30, 17, PURPLE, True, HEAD, PP_ALIGN.CENTER)
text(p, "UNDO EACH PITCH BY ONE STEP", 4.24, 3.52, 4.92, 0.34, 13, PURPLE,
     True, HEAD, PP_ALIGN.CENTER)
label(p, "ORIGINAL PITCHES", 6.86, 2.30, 3.3, fill=INK)
token_row(p, ["C", "E", "G", "A", "D"], 6.86, 2.69, 5.65, LAVENDER, 23)
rule(p, 0.82, 3.91, 11.71, INK, 0.018)
label(p, "PLACE THE FIVE ORIGINAL PITCHES IN THE FIVE NOTE POSITIONS", 0.82, 4.22,
      10.9, fill=INK, size=10)
slot_board(p, ["C", "REST", "E", "G", "REST", "A", "D", "REST"],
           0.82, 4.62, 11.69, 1.31, captions=True, font_size=29)
text(p, "Slots 2, 5 and 8 are rests. The note sequence stays C → E → G → A → D.",
     0.83, 6.30, 11.66, 0.43, 16, MUTED)

# 09 — Communication and shared board
p = slide(9, "Talk, place, and listen.",
          "The puzzle is cooperative: clue buttons and chat help the Conductor fill the board.")
box(p, 0.80, 2.36, 5.59, 3.72, WHITE, INK, 1.2)
label(p, "THE CONVERSATION", 1.04, 2.62, 4.84, fill=INK, size=11)
for y, bg, who, message in (
    (3.05, LIME_SOFT, "RHYTHM", "Slots 2, 5 and 8 are rests."),
    (3.83, PEACH_SOFT, "MELODIST", "Shifted notes: D, G, A, C, E."),
    (4.61, LAVENDER_SOFT, "CONDUCTOR", "One step back. I'll fill the board."),
):
    box(p, 1.04, y, 5.09, 0.64, bg, RULE, 0.8)
    label(p, who, 1.18, y + 0.10, 1.27, 0.16, INK, 8)
    text(p, message, 1.18, y + 0.30, 4.67, 0.27, 13.5, INK)
text(p, "Share buttons send one role-specific clue; chat allows questions.",
     1.04, 5.53, 5.05, 0.42, 13, MUTED)
box(p, 6.65, 2.36, 5.87, 3.72, LAVENDER_SOFT, INK, 1.2)
label(p, "THE SHARED BOARD", 6.89, 2.62, 4.84, fill=INK, size=11)
slot_board(p, ["C", "R", "E", "?", "R", "A", "D", "R"],
           6.89, 3.11, 5.39, 1.16, numbers=True, font_size=19)
text(p, "Only the Conductor can select a slot and choose a pitch, Rest, or Clear.",
     6.89, 4.56, 5.31, 0.70, 16, INK, True, HEAD)
text(p, "All three players see board updates immediately.", 6.89, 5.47,
     5.20, 0.37, 14, MUTED)
box(p, 0.80, 6.33, 11.73, 0.48, LIME, INK, 1.0)
text(p, "PLAY = hear the current board locally in your browser. No voice call is needed.",
     1.02, 6.42, 11.22, 0.26, 13, INK, True, BODY)

# 10 — Submit and score
p = slide(10, "Finish, confirm, then score.",
          "The song is checked only after the Conductor confirms the complete answer.")
for i, (title, detail, bg) in enumerate((
    ("FILL ALL 8", "Every slot needs C, D, E, G, A or an explicit Rest.", WHITE),
    ("SUBMIT", "All players see the proposed board before it is locked.", LAVENDER),
    ("CONFIRM", "The Conductor confirms; the server compares it to the target.", LIME),
)):
    x = 0.80 + i * 4.04
    box(p, x, 2.36, 3.77, 1.65, bg, INK, 1.2)
    label(p, f"0{i + 1}  {title}", x + 0.19, 2.60, 3.38, fill=INK, size=12)
    text(p, detail, x + 0.19, 2.99, 3.34, 0.82, 15, INK)
rule(p, 0.81, 4.36, 11.70, INK, 0.018)
text(p, "1 point", 0.82, 4.72, 2.30, 0.60, 30, INK, True, HEAD)
text(p, "for each pitch or Rest in the correct slot", 3.03, 4.85,
     6.74, 0.39, 20, INK, True, HEAD)
box(p, 10.38, 4.59, 2.10, 0.88, LIME, INK, 1.25)
text(p, "MAX 8/8", 10.48, 4.77, 1.90, 0.42, 20, INK, True, HEAD,
     PP_ALIGN.CENTER)
text(p, "Wrong pitch, misplaced note, or missing rest = 0 for that slot.",
     0.83, 5.78, 11.59, 0.39, 17, MUTED)
text(p, "A question mark is not a Rest; unanswered boards cannot be submitted.",
     0.83, 6.31, 11.59, 0.39, 17, MUTED)

# 11 — Reveal and next round
p = slide(11, "The reveal shows what changed.",
          "Everyone sees the submitted song, the true target, and an exact comparison.")
box(p, 0.80, 2.40, 2.28, 2.42, LAVENDER, INK, 1.3)
label(p, "EXAMPLE SCORE", 1.03, 2.71, 1.80, fill=INK, size=10)
text(p, "7", 1.01, 3.00, 1.24, 1.18, 66, INK, True, HEAD)
text(p, "/ 8", 2.09, 3.60, 0.78, 0.45, 22, INK, True, HEAD)
text(p, "One slot was wrong.", 1.04, 4.35, 1.77, 0.37, 13, MUTED)
label(p, "YOUR VERSION", 3.45, 2.46, 2.72, fill=INK, size=10)
slot_board(p, ["C", "R", "E", "A", "R", "A", "D", "R"],
           3.45, 2.83, 9.06, 0.88, numbers=False, font_size=19,
           miss_indices={3})
label(p, "CORRECT TARGET", 3.45, 3.92, 2.72, fill=INK, size=10)
slot_board(p, ["C", "R", "E", "G", "R", "A", "D", "R"],
           3.45, 4.29, 9.06, 0.88, numbers=False, font_size=19)
box(p, 6.86, 5.36, 1.10, 0.34, MISS, INK, 0.8)
text(p, "Slot 4: A should be G", 8.12, 5.36, 3.05, 0.35, 14,
     INK, True, HEAD)
pill(p, "PLAY YOUR VERSION", 0.80, 6.06, 2.54, WHITE, 0.47, 12)
pill(p, "PLAY TARGET", 3.49, 6.06, 1.91, WHITE, 0.47, 12)
pill(p, "REPLAY SAME PUZZLE", 6.39, 6.06, 2.60, LAVENDER, 0.47, 12)
pill(p, "NEW PUZZLE", 9.14, 6.06, 2.13, LIME, 0.47, 12)
text(p, "The host chooses the next round; everyone readies again.", 0.82,
     6.70, 11.59, 0.27, 13, MUTED)

# 12 — Recovery and invitation
p = slide(12, "Ready to play?",
          "A quick safety net, plus an easy way to learn before inviting friends.")
shadow_card(p, 0.80, 2.38, 5.56, 3.22, LAVENDER_SOFT, LAVENDER)
label(p, "IF A PLAYER DISCONNECTS", 1.06, 2.66, 4.93, fill=INK, size=11)
text(p, "The room pauses for 60 seconds.", 1.06, 3.11, 4.96, 0.56,
     22, INK, True, HEAD)
text(p, "Reopen the game in the same browser to reclaim your role, clue and board progress.",
     1.06, 3.82, 4.96, 0.94, 16, MUTED)
text(p, "If time runs out, the host can end the round.", 1.06, 5.00,
     4.90, 0.42, 14, INK)
shadow_card(p, 6.72, 2.38, 5.80, 3.22, LIME_SOFT, LIME)
label(p, "PRACTICE ALONE", 6.98, 2.66, 5.10, fill=INK, size=11)
text(p, "Choose Try Demo.", 6.98, 3.11, 5.08, 0.56,
     22, INK, True, HEAD)
text(p, "Switch between the three roles, inspect their clues and finish a practice puzzle.",
     6.98, 3.82, 5.10, 0.94, 16, MUTED)
text(p, "Then create a live room and invite two players.", 6.98, 5.00,
     5.08, 0.42, 14, INK)
box(p, 0.80, 6.13, 11.72, 0.61, INK, None)
text(p, "THREE CLUES  →  ONE SHARED BOARD  →  EIGHT EXACT SLOTS",
     1.13, 6.26, 11.04, 0.30, 17, LIME, True, HEAD, PP_ALIGN.CENTER)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(f"Created {OUTPUT} ({len(prs.slides)} slides)")
